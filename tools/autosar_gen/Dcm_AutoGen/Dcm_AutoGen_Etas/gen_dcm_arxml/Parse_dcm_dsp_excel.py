"""解析 UDS 诊断问卷 Excel 的 DcmDspSession / DcmDspSecurity 表 (DcmDsp 行级配置).

与 Service List 不同: 两表首行即表头 (无标题行), 列名 = arxml 参数定义名,
空单元格 = 空属性 (不生成对应参数).
"""

import sys
from types import SimpleNamespace

import openpyxl

from Rebuild_DcmDsp import SESSION_CONTAINER, SECURITY_CONTAINER, \
    SESSION_ROW_PARAMS, SECURITY_ROW_PARAMS

SESSION_SHEET = SESSION_CONTAINER
SECURITY_SHEET = SECURITY_CONTAINER
SESSION_NAME_COL = "DcmDspSessionRow"
SECURITY_NAME_COL = "DcmDspSecurityRow"

# Excel 中存在、但工具暂不生成的列 (基线 arxml 无对应参数, 非空时报错提醒扩展)
SECURITY_UNSUPPORTED_COLS = ["DcmRbDspSecurityRowUseAsynchronousServerCallPoint",
                             "DcmRbDspSecurityRowUseFlexibleLength"]


def _fail(msg):
    raise ValueError("DcmDsp 表解析错误: %s" % msg)


def _norm_name(name):
    """SHORT-NAME 规范化: 去首尾空白, 空格转下划线."""
    return str(name).strip().replace(" ", "_")


def _to_int(v, what):
    if isinstance(v, bool):
        _fail("%s 必须是整数, 实际: %r" % (what, v))
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        if v != int(v):
            _fail("%s 必须是整数, 实际: %r" % (what, v))
        return int(v)
    try:
        return int(str(v).strip(), 10)
    except ValueError:
        _fail("%s 必须是整数, 实际: %r" % (what, v))


def _to_float(v, what):
    if isinstance(v, bool):
        _fail("%s 必须是数值, 实际: %r" % (what, v))
    try:
        return float(v)
    except (TypeError, ValueError):
        _fail("%s 必须是数值, 实际: %r" % (what, v))


def _to_bool(v, what):
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    if s == "true":
        return True
    if s == "false":
        return False
    _fail("%s 必须是 True/False, 实际: %r" % (what, v))


def _convert(kind, v, what):
    if kind == "INTEGER":
        return _to_int(v, what)
    if kind == "FLOAT":
        return _to_float(v, what)
    if kind == "BOOLEAN":
        return _to_bool(v, what)
    return str(v).strip()


def _parse_sheet(wb, sheet_name, name_col, row_params, unsupported_cols=None):
    """解析单个 DcmDsp 行表, 返回 [{"name": SHORT-NAME, <参数定义名>: 类型化值}, ...].

    空单元格的参数不出现在结果中 (空属性不生成).
    """
    if sheet_name not in wb.sheetnames:
        _fail("缺少工作表 %r" % sheet_name)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        _fail("工作表 %r 为空" % sheet_name)

    # 首行即表头 (与 Service List 的标题行+表头不同)
    header = rows[0]
    col_idx = {}
    for i, h in enumerate(header):
        if h is not None:
            col_idx[str(h).strip()] = i
    required = [name_col] + [p for p, _ in row_params]
    missing = [c for c in required if c not in col_idx]
    if missing:
        _fail("%r 缺少必需列: %s" % (sheet_name, ", ".join(missing)))

    out = []
    names = set()
    for ridx, row in enumerate(rows[1:], start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        what = "%r 第 %d 行" % (sheet_name, ridx)

        name_v = row[col_idx[name_col]]
        if name_v is None or str(name_v).strip() == "":
            _fail("%s 的 %s 列为空" % (what, name_col))
        name = _norm_name(name_v)
        if name in names:
            _fail("%s 行名重复: %s" % (what, name))
        names.add(name)

        if unsupported_cols:
            for col in unsupported_cols:
                if col in col_idx:
                    v = row[col_idx[col]]
                    if v is not None and str(v).strip() != "":
                        _fail("%s 的 %s 列暂不支持生成 (非空值 %r), 请先扩展工具"
                              % (what, col, v))

        rec = {"name": name}
        for param, kind in row_params:
            v = row[col_idx[param]]
            if v is None or str(v).strip() == "":
                continue
            rec[param] = _convert(kind, v, "%s 的 %s" % (what, param))
        if len(rec) == 1:
            _fail("%s 除行名外全部参数为空" % what)
        out.append(rec)

    if not out:
        _fail("工作表 %r 未解析到任何数据行" % sheet_name)
    return out


def parse_dcm_dsp_excel(excel_path) -> SimpleNamespace:
    """
    返回 SimpleNamespace(session_rows, security_rows, session_count, security_count).
    行 = {"name": SHORT-NAME, <参数定义名>: 类型化值}; 空属性不出现.
    """
    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    except Exception as e:
        _fail("打开 Excel 失败 %s: %s" % (excel_path, e))
    session_rows = _parse_sheet(wb, SESSION_SHEET, SESSION_NAME_COL, SESSION_ROW_PARAMS)
    security_rows = _parse_sheet(wb, SECURITY_SHEET, SECURITY_NAME_COL,
                                 SECURITY_ROW_PARAMS, SECURITY_UNSUPPORTED_COLS)
    wb.close()
    return SimpleNamespace(session_rows=session_rows, security_rows=security_rows,
                           session_count=len(session_rows),
                           security_count=len(security_rows))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("用法: python Parse_dcm_dsp_excel.py <excel_path>")
        sys.exit(1)
    try:
        data = parse_dcm_dsp_excel(sys.argv[1])
    except ValueError as e:
        print("错误：%s" % e)
        sys.exit(1)
    print("解析成功: %d 个 Session 行, %d 个 Security 行"
          % (data.session_count, data.security_count))
    for r in data.session_rows:
        print("  [Session]  %-28s %s" % (r["name"], r))
    for r in data.security_rows:
        print("  [Security] %-28s %s" % (r["name"], r))
