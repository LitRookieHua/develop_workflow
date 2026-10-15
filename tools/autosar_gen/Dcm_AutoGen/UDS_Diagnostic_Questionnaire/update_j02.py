# -*- coding: utf-8 -*-
"""更新模板与 J02:

1. Service List: NRC 后新增 NRC22 列(Y/N),模板+J02 同步
2. DID List: Locked/App/Boot/iMotion Unlocked 改为 22:xx 2E:xx 格式(模板+J02);
   J02 修正: F187->VersionDid; F18C 支持扩展+iMotion Unlocked 写入; F1DF CAN固定回复0x00
3. RID List(J02): 仅保留 8510/8511, 其余为示例删除
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList

TPL = 'UDS_Diagnostic_Questionnaire_Template.xlsx'
J02 = 'J02/UDS_Diagnostic_Questionnaire_J02.xlsx'

HEAD_FILL = PatternFill('solid', fgColor='1F4E78')
HEAD_FONT = Font(name='宋体', bold=True, color='FFFFFF', size=10)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
THIN = Side(style='thin', color='B0B0B0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
COMBO = '"22:Yes 2E:No,22:Yes 2E:Yes,22:No 2E:No,22:No 2E:Yes"'

SVC_HEADERS = ['Service (Hex)', 'Name', 'Sub-function (Hex（Name）)', 'NRC',
               'NRC22', 'SPRMIB', 'Addressing Mode', 'Default Session',
               'Programming Session', 'Extended Session', 'Locked',
               'App Unlocked', 'Boot Unlocked', 'iMotion Unlocked', 'JIRA',
               'Remark']
SVC_WIDTHS = [10, 26, 24, 16, 8, 8, 10, 10, 12, 12, 10, 12, 12, 12, 12, 30]


def add_dv(ws, formula, cells):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True,
                        errorStyle='warning')
    ws.add_data_validation(dv)
    dv.add(cells)


def rebuild_service(ws, data):
    """data: list of 16-col row lists, or None (模板空)"""
    ws.data_validations = DataValidationList()
    ncols = len(SVC_HEADERS)
    for idx, h in enumerate(SVC_HEADERS, 1):
        c = ws.cell(row=2, column=idx, value=h)
        c.fill = HEAD_FILL; c.font = HEAD_FONT; c.alignment = CENTER; c.border = BORDER
    for cidx in range(ncols + 1, ws.max_column + 1):
        c = ws.cell(row=2, column=cidx)
        c.value = None; c.fill = PatternFill(); c.font = Font(); c.border = Border()
    ws.row_dimensions[2].height = 30
    for i, w in enumerate(SVC_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col in list(ws.column_dimensions.keys()):
        if openpyxl.utils.column_index_from_string(col) > ncols:
            del ws.column_dimensions[col]
    for r in range(3, 33):
        for cidx in range(1, ncols + 1):
            cell = ws.cell(row=r, column=cidx)
            cell.border = BORDER
            if data and r - 3 < len(data):
                v = data[r - 3][cidx - 1]
                if v is not None and v != '':
                    cell.value = v
    ws.freeze_panes = 'D3'
    ws.auto_filter.ref = f'A2:P32'
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row == 1:
            ws.unmerge_cells(str(mr))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    add_dv(ws, '"Y,N"', 'F3:F32')
    add_dv(ws, '"P,F,P/F"', 'G3:G32')
    add_dv(ws, '"Y,N"', 'H3:N32')
    add_dv(ws, '"Y,N"', 'E3:E32')  # NRC22


def rebuild_did_dv(ws):
    ws.data_validations = DataValidationList()
    add_dv(ws, '"Yes,No"', 'C3:D32')
    add_dv(ws, '"Hex,ASCII,BCD"', 'E3:E32')
    add_dv(ws, '"NvmDid,VersionDid,DynamicDid"', 'F3:F32')
    add_dv(ws, COMBO, 'I3:K32')   # 会话
    add_dv(ws, COMBO, 'L3:O32')   # 解锁等级


# ============ 1. 模板 ============
wb = openpyxl.load_workbook(TPL)
rebuild_service(wb['Service List'], None)
rebuild_did_dv(wb['DID List'])
wb.save(TPL)
print('template updated')

# ============ 2. J02 ============
wb = openpyxl.load_workbook(J02)

# ---- Service List: 读旧数据, 插入 NRC22 ----
NRC22 = {'0x10': 'Y', '0x11': 'Y', '0x14': 'Y', '0x19': 'N', '0x22': 'Y',
         '0x27': 'Y', '0x28': 'Y', '0x2E': 'Y', '0x31': 'Y', '0x36': 'N',
         '0x37': 'N', '0x38': 'Y', '0x3E': 'N', '0x85': 'Y'}
ws = wb['Service List']
old = []
for r in range(3, 33):
    row = [ws.cell(row=r, column=c).value for c in range(1, 16)]
    if any(v is not None for v in row):
        old.append(row)
new = []
for row in old:
    sid = row[0]
    new.append(row[:4] + [NRC22.get(sid, '')] + row[4:])
rebuild_service(ws, new)
print(f'J02 Service List: {len(new)} rows with NRC22')

# ---- DID List: 解锁列 22:2E 格式 + 修正 ----
ws = wb['DID List']
for r in range(3, 33):
    did = ws.cell(row=r, column=1).value
    if not did:
        continue
    # 默认: 解锁列全部 22:Yes 2E:No
    for c in (12, 13, 14, 15):
        ws.cell(row=r, column=c).value = '22:Yes 2E:No'
    if did == '0xF15A':   # 无$22, 仅Boot(√2)$2E
        for c in (12, 13, 14, 15):
            ws.cell(row=r, column=c).value = '22:No 2E:No'
        ws.cell(row=r, column=14).value = '22:No 2E:Yes'
    elif did == '0xF190':  # $2E √1 App(扩展会话)
        ws.cell(row=r, column=13).value = '22:Yes 2E:Yes'
    elif did == '0x2EE2':  # $2E √1 App(扩展会话)
        ws.cell(row=r, column=13).value = '22:Yes 2E:Yes'
    elif did == '0xF18C':  # 用户: 支持扩展+iMotion Unlocked 写入
        ws.cell(row=r, column=4).value = 'Yes'
        ws.cell(row=r, column=11).value = '22:Yes 2E:Yes'
        ws.cell(row=r, column=15).value = '22:Yes 2E:Yes'
        ws.cell(row=r, column=17).value = ('结构: SupplierCode(3B)+PartCode(9B)+Batch(1B)+PartNo(4B)+Date(3B); '
                                           '供应商代码待客户提供; $2E仅扩展会话iMotion Unlocked')
    elif did == '0xF187':  # 用户: VersionDid
        ws.cell(row=r, column=6).value = 'VersionDid'
    elif did == '0xF1DF':  # 用户: CAN固定回复0x00
        ws.cell(row=r, column=7).value = '0x00'
        ws.cell(row=r, column=17).value = ('bit7-4 Reprogrammability; bit3-0 ProgramConsistency; '
                                           'CAN固定回复0x00')
rebuild_did_dv(ws)
print('J02 DID List: unlock columns -> 22:2E format, F187/F18C/F1DF fixed')

# ---- RID List: 仅保留 8510/8511 ----
ws = wb['RID List']
keep = []
for r in range(3, 33):
    rid = ws.cell(row=r, column=1).value
    if rid in ('0x8510', '0x8511'):
        keep.append([ws.cell(row=r, column=c).value for c in range(1, 15)])
for r in range(3, 33):
    for c in range(1, 15):
        ws.cell(row=r, column=c).value = None
for i, row in enumerate(keep):
    for cidx, v in enumerate(row, 1):
        if v is not None:
            ws.cell(row=3 + i, column=cidx).value = v
print(f'J02 RID List: kept {len(keep)} rows (8510/8511 only)')

# ---- History: 更新说明 ----
ws = wb['History']
ws['D3'] = ('Initialized from FAWVW_UDS_Diagnostic_Questionnaire_MFK_DK3_J02_V001_260920: '
            'General(48) + Service(14 svc/29 rows, NRC22) + DID(22, unlock 22:2E format) + '
            'RID(2 RIDs/4 rows: 8510/8511, 其余为客户示例) + NVM(5)')
print('J02 History updated')

wb.save(J02)
print('saved:', J02)
