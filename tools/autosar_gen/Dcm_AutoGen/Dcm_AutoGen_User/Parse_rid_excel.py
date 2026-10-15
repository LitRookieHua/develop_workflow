"""解析 UDS 诊断问卷的 RID List 表 (供 gen_rid_code 使用).

RID List 每行描述一个 (RID, Control Type) 组合: 同一 RID 拆成多行,
Control Type 形如 "0x01（StartRoutine）" / "0x03（RequestRoutineResults）".
权限列为 7 列独立的 'Y'/'N' (三会话 + Locked/App/Boot/iMotion Unlocked).

用法:
  python Parse_rid_excel.py <excel_path>
"""

import re
import sys
from types import SimpleNamespace

import openpyxl

RID_SHEET = "RID List"

RID_REQUIRED_COLS = [
    "RID (Hex)", "Name", "Control Type", "Request Data", "Request Data Size",
    "Response Data", "Response Data Size",
    "Default Session", "Programming Session", "Extended Session",
    "Locked", "App Unlocked", "Boot Unlocked", "iMotion Unlocked",
    "JIRA", "Remark",
]

CONTROL_RE = re.compile(r"^0x([0-9A-Fa-f]{2})\b")
VALID_CONTROL = {0x01: "StartRoutine", 0x02: "StopRoutine", 0x03: "RequestRoutineResults"}
SKIP_REMARK_KEYWORD = "Docan不实现"

SESSION_COLS = [("Default Session", "DEFAULT"),
                ("Programming Session", "PROGRAMMING"),
                ("Extended Session", "EXTENDED")]
SEC_COLS = [("Locked", "LOCK"),
            ("App Unlocked", "APP"),
            ("Boot Unlocked", "BOOT"),
            ("iMotion Unlocked", "IMOTION")]


def _fail(msg):
    raise ValueError("RID 问卷解析错误: %s" % msg)


def _parse_hex(v, what):
    s = str(v).strip()
    try:
        return int(s, 16)
    except ValueError:
        _fail("%s 非法十六进制值: %r" % (what, s))


def _cell_yn(row, col_idx, header, what):
    v = row[col_idx[header]]
    s = str(v).strip() if v is not None else ""
    if s not in ("Y", "N"):
        _fail("%s 的 %s 列必须为 Y/N, 实际: %r" % (what, header, v))
    return s == "Y"


def _cell_size(row, col_idx, header, what):
    v = row[col_idx[header]]
    try:
        n = int(v)
    except (TypeError, ValueError):
        _fail("%s 的 %s 非法: %r" % (what, header, v))
    if n < 0:
        _fail("%s 的 %s 必须为非负整数: %d" % (what, header, n))
    return n


def _header_index(rows, sheet):
    if len(rows) < 3:
        _fail("%s 工作表为空" % sheet)
    col_idx = {}
    for i, h in enumerate(rows[1]):
        if h is not None:
            col_idx[str(h).strip()] = i
    return col_idx, len(rows[1])


def _iter_data_rows(ws, sheet):
    """返回 (行号, 补齐宽度的值列表); 第 1 行标题, 第 2 行表头, 第 3 行起数据."""
    rows = list(ws.iter_rows(values_only=True))
    col_idx, width = _header_index(rows, sheet)
    for ridx, row in enumerate(rows[2:], start=3):
        if row is None:
            continue
        vals = list(row) + [None] * (width - len(row))
        yield ridx, vals, col_idx


def parse_rid_excel(excel_path) -> SimpleNamespace:
    """
    返回 SimpleNamespace(rids=[...]):
    rids: [{rid, name, controls: {ctrl: {"req_size", "resp_size"}},
            sessions: set, sec: set}]   (RID 升序, ctrl 为 0x01/0x02/0x03)
    """
    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    except Exception as e:
        _fail("打开 Excel 失败 %s: %s" % (excel_path, e))

    if RID_SHEET not in wb.sheetnames:
        _fail("缺少工作表 %r" % RID_SHEET)

    rids = []
    rid_by_id = {}
    skipped = []

    for ridx, row, col_idx in _iter_data_rows(wb[RID_SHEET], RID_SHEET):
        rid_hex = row[col_idx["RID (Hex)"]]
        if rid_hex is None or str(rid_hex).strip() == "":
            continue
        what = "第 %d 行" % ridx

        remark = row[col_idx["Remark"]]
        if remark is not None and SKIP_REMARK_KEYWORD in str(remark):
            skipped.append(str(rid_hex).strip())
            continue

        rid = _parse_hex(rid_hex, what + " RID (Hex)")

        ctrl_v = row[col_idx["Control Type"]]
        ctrl_s = str(ctrl_v).strip() if ctrl_v is not None else ""
        m = CONTROL_RE.match(ctrl_s)
        if not m:
            _fail("0x%04X 的 Control Type 必须形如 '0x01（StartRoutine）', 实际: %r" % (rid, ctrl_v))
        ctrl = int(m.group(1), 16)
        if ctrl not in VALID_CONTROL:
            _fail("0x%04X 的 Control Type 子功能非法: 0x%02X" % (rid, ctrl))

        name_v = row[col_idx["Name"]]
        name = str(name_v).strip() if name_v is not None else ""
        if not name:
            _fail("0x%04X 的 Name 为空" % rid)

        req_size = _cell_size(row, col_idx, "Request Data Size", what)
        resp_size = _cell_size(row, col_idx, "Response Data Size", what)

        sessions = set()
        for header, part in SESSION_COLS:
            if _cell_yn(row, col_idx, header, what):
                sessions.add(part)
        sec = set()
        for header, part in SEC_COLS:
            if _cell_yn(row, col_idx, header, what):
                sec.add(part)

        entry = rid_by_id.get(rid)
        if entry is None:
            entry = {"rid": rid, "name": name, "controls": {},
                     "sessions": sessions, "sec": sec}
            rids.append(entry)
            rid_by_id[rid] = entry
        else:
            if entry["name"] != name:
                _fail("0x%04X 的 Name 不一致: %r vs %r" % (rid, entry["name"], name))
            if entry["sessions"] != sessions or entry["sec"] != sec:
                _fail("0x%04X 的权限列在不同 Control Type 行之间不一致" % rid)
        if ctrl in entry["controls"]:
            _fail("0x%04X 的 Control Type 0x%02X 重复" % (rid, ctrl))
        entry["controls"][ctrl] = {"req_size": req_size, "resp_size": resp_size}

    wb.close()

    if skipped:
        print("已跳过 Remark 含 %r 的 RID: %s" % (SKIP_REMARK_KEYWORD, ", ".join(skipped)))
    if not rids:
        _fail("未解析到任何 RID")

    rids.sort(key=lambda x: x["rid"])
    return SimpleNamespace(rids=rids)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("用法: python Parse_rid_excel.py <excel_path>")
        sys.exit(1)
    try:
        data = parse_rid_excel(sys.argv[1])
    except ValueError as e:
        print("错误：%s" % e)
        sys.exit(1)
    print("解析成功: %d 个 RID" % len(data.rids))
    for r in data.rids:
        ss = "/".join(sorted(r["sessions"])) or "-"
        sc = "/".join(sorted(r["sec"])) or "-"
        print("  0x%04X %-30s sessions[%s] sec[%s]" % (r["rid"], r["name"], ss, sc))
        for ctrl in sorted(r["controls"]):
            c = r["controls"][ctrl]
            print("    0x%02X %-22s req=%dB resp=%dB"
                  % (ctrl, VALID_CONTROL[ctrl], c["req_size"], c["resp_size"]))
