import xml.etree.ElementTree as ET
from xml.dom import minidom
import openpyxl
import os
import copy
import argparse
import sys

# AUTOSAR Namespace
NS = {'ns': 'http://autosar.org/schema/r4.0'}

def load_excel_data(file_path):
    """从 Excel 读取数据"""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Excel 文件不存在: {file_path}")
        
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    sheet_name = None
    for name in wb.sheetnames:
        if '01_nvm_list' in name.lower():
            sheet_name = name
            break
    
    if not sheet_name:
        for name in wb.sheetnames:
            if 'nvm' in name.lower() and 'list' in name.lower():
                sheet_name = name
                break
    
    if not sheet_name:
        raise ValueError("未找到包含 'NVM_List' 的工作表")
    
    ws = wb[sheet_name]
    data_map = {}
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] is None: 
            continue
            
        block_name_original = str(row[1]).strip()
        try:
            block_length = int(row[2])
        except (ValueError, TypeError):
            print(f"警告: 跳过行，长度无效: {row}")
            continue
        
        # 确保 Key 带 _Ram 后缀
        if not block_name_original.endswith('_Ram'):
            block_name_key = f"{block_name_original}_Ram"
        else:
            block_name_key = block_name_original
            
        data_map[block_name_key] = {
            'length': block_length,
            'name_original': block_name_original,
            'name_final': block_name_key
        }
        
    return data_map

def get_template_container(root):
    """查找一个现有的 Pdm_Z..._Ram 块作为模板"""
    for container in root.iter('{http://autosar.org/schema/r4.0}ECUC-CONTAINER-VALUE'):
        def_ref_elem = container.find('ns:DEFINITION-REF', NS)
        if def_ref_elem is None or 'NvMBlockDescriptor' not in def_ref_elem.text:
            continue
            
        for param_val in container.findall('ns:PARAMETER-VALUES/ns:ECUC-TEXTUAL-PARAM-VALUE', NS):
            def_ref = param_val.find('ns:DEFINITION-REF', NS)
            if def_ref is not None and 'NvMRamBlockDataAddress' in def_ref.text:
                val_elem = param_val.find('ns:VALUE', NS)
                if val_elem is not None:
                    val_text = val_elem.text.strip()
                    if val_text.startswith('Pdm_Z') and val_text.endswith('_Ram'):
                        return container
    return None

def create_new_block(template_container, ram_address_name_final, length, ram_address_name_original):
    """基于模板创建新块"""
    if template_container is None:
        raise Exception("无法在 XML 中找到 Pdm_Z 类型的模板块来创建新块。")

    new_container = copy.deepcopy(template_container)
    
    suffix = ram_address_name_original.replace('Pdm_', '', 1)
    new_short_name = f"NVM_ID_{suffix}"
    
    short_name_elem = new_container.find('ns:SHORT-NAME', NS)
    if short_name_elem is not None:
        short_name_elem.text = new_short_name
    else:
        sn_elem = ET.SubElement(new_container, '{http://autosar.org/schema/r4.0}SHORT-NAME')
        sn_elem.text = new_short_name

    for param_val in new_container.findall('ns:PARAMETER-VALUES/ns:ECUC-TEXTUAL-PARAM-VALUE', NS):
        def_ref = param_val.find('ns:DEFINITION-REF', NS)
        if def_ref is not None and 'NvMRamBlockDataAddress' in def_ref.text:
            val_elem = param_val.find('ns:VALUE', NS)
            if val_elem is not None:
                val_elem.text = ram_address_name_final
                break
    
    for param_val in new_container.findall('ns:PARAMETER-VALUES/ns:ECUC-NUMERICAL-PARAM-VALUE', NS):
        def_ref = param_val.find('ns:DEFINITION-REF', NS)
        if def_ref is not None and 'NvMNvBlockLength' in def_ref.text:
            val_elem = param_val.find('ns:VALUE', NS)
            if val_elem is not None:
                val_elem.text = str(length)
                break
    
    return new_container

def prettify_xml(elem):
    """
    使用 minidom 对 XML 进行美化（添加缩进和换行）
    返回不含 XML 声明的内容字符串
    """
    rough_string = ET.tostring(elem, encoding='utf-8')
    reparsed = minidom.parseString(rough_string)
    # indent="  " 使用2个空格缩进，可根据需要调整
    pretty = reparsed.toprettyxml(indent="  ", encoding='utf-8').decode('utf-8')
    
    lines = pretty.split('\n')
    cleaned_lines = []
    skip_decl = True
    
    for line in lines:
        # 跳过 minidom 生成的 XML 声明行
        if skip_decl and line.strip().startswith('<?xml'):
            skip_decl = False
            continue
        # 移除纯空行（minidom 经常会在标签间插入多余空行）
        if line.strip() == '':
            continue
        cleaned_lines.append(line)
    
    return '\n'.join(cleaned_lines)

def update_xml(xml_path, excel_data, output_path):
    # 1. 读取并保存原始文件的第一行 (XML Declaration)
    original_header = ""
    try:
        with open(xml_path, 'r', encoding='utf-8') as f:
            first_line = f.readline()
            if first_line.strip().startswith('<?xml'):
                original_header = first_line.strip() + '\n'
            else:
                original_header = '<?xml version="1.0" encoding="UTF-8"?>\n'
    except Exception as e:
        print(f"警告: 读取文件头失败。错误: {e}")
        original_header = '<?xml version="1.0" encoding="UTF-8"?>\n'

    ET.register_namespace('', 'http://autosar.org/schema/r4.0')
    ET.register_namespace('xsi', 'http://www.w3.org/2001/XMLSchema-instance')
    
    tree = ET.parse(xml_path)
    root = tree.getroot()
    
    containers_parent = None
    for elem in root.iter('{http://autosar.org/schema/r4.0}ECUC-MODULE-CONFIGURATION-VALUES'):
        sn = elem.find('ns:SHORT-NAME', NS)
        if sn is not None and sn.text == 'NvM':
            containers_parent = elem.find('ns:CONTAINERS', NS)
            break
            
    if containers_parent is None:
        for elem in root.iter('{http://autosar.org/schema/r4.0}CONTAINERS'):
             containers_parent = elem
             break

    if containers_parent is None:
        raise Exception("未在 XML 中找到 CONTAINERS 节点")

    blocks_to_remove = []
    existing_ram_addresses = set()
    template_block = None
    
    # --- 第一阶段：扫描、更新、标记删除 ---
    for container in list(containers_parent):
        def_ref_elem = container.find('ns:DEFINITION-REF', NS)
        if def_ref_elem is None or 'NvMBlockDescriptor' not in def_ref_elem.text:
            continue
            
        ram_address_name = None
        for param_val in container.findall('ns:PARAMETER-VALUES/ns:ECUC-TEXTUAL-PARAM-VALUE', NS):
            def_ref = param_val.find('ns:DEFINITION-REF', NS)
            if def_ref is not None and 'NvMRamBlockDataAddress' in def_ref.text:
                val_elem = param_val.find('ns:VALUE', NS)
                if val_elem is not None:
                    ram_address_name = val_elem.text.strip()
                break
        
        if ram_address_name:
            existing_ram_addresses.add(ram_address_name)
            
            if ram_address_name.startswith('Pdm_Z') and ram_address_name.endswith('_Ram'):
                if template_block is None:
                    template_block = container
                
                if ram_address_name not in excel_data:
                    blocks_to_remove.append(container)
                    print(f"[删除] Block 地址: {ram_address_name}")
                else:
                    new_length = excel_data[ram_address_name]['length']
                    for param_val in container.findall('ns:PARAMETER-VALUES/ns:ECUC-NUMERICAL-PARAM-VALUE', NS):
                        def_ref = param_val.find('ns:DEFINITION-REF', NS)
                        if def_ref is not None and 'NvMNvBlockLength' in def_ref.text:
                            val_elem = param_val.find('ns:VALUE', NS)
                            if val_elem is not None:
                                old_val = val_elem.text
                                val_elem.text = str(new_length)
                                if old_val != str(new_length):
                                    print(f"[更新] Block 地址: {ram_address_name} | 长度: {old_val} -> {new_length}")
                            break

    for container in blocks_to_remove:
        containers_parent.remove(container)

    # --- 第二阶段：添加新 Block ---
    if template_block is None:
        for container in containers_parent:
             def_ref_elem = container.find('ns:DEFINITION-REF', NS)
             if def_ref_elem is None or 'NvMBlockDescriptor' not in def_ref_elem.text:
                 continue
             for param_val in container.findall('ns:PARAMETER-VALUES/ns:ECUC-TEXTUAL-PARAM-VALUE', NS):
                def_ref = param_val.find('ns:DEFINITION-REF', NS)
                if def_ref is not None and 'NvMRamBlockDataAddress' in def_ref.text:
                    val_elem = param_val.find('ns:VALUE', NS)
                    if val_elem is not None:
                        v = val_elem.text.strip()
                        if v.startswith('Pdm_Z') and v.endswith('_Ram'):
                            template_block = container
                            break
             if template_block: break

    added_count = 0
    for ram_addr_key, info in excel_data.items():
        if ram_addr_key.startswith('Pdm_Z') and ram_addr_key not in existing_ram_addresses:
            if template_block is None:
                print(f"[错误] 需要添加新块 {ram_addr_key}，但未找到可用的模板块。跳过。")
                continue
            
            print(f"[新增] Block 地址: {ram_addr_key}, 长度: {info['length']}")
            new_block = create_new_block(
                template_block, 
                ram_addr_key, 
                info['length'], 
                info['name_original']
            )
            containers_parent.append(new_block)
            added_count += 1

    # --- 第三阶段：保存文件 ---
    
    # 1. 获取美化后的 XML 内容（不含声明）
    xml_content = prettify_xml(root)
    
    # 2. 拼接：原始头部 + 美化内容
    # 注意：prettify_xml 返回的字符串末尾没有换行，我们需要在这里控制
    final_output = original_header + xml_content
    
    # 3. 【关键步骤】强制在文件末尾添加两个换行符，以保留两行空行
    # 这样文件结束时会呈现为：...</AUTOSAR>\n\n (EOF)
    final_output += "\n\n"
    
    # 4. 写入文件
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(final_output)
    
    print(f"\n处理完成。")
    print(f"删除了 {len(blocks_to_remove)} 个旧块。")
    print(f"新增了 {added_count} 个新块。")
    print(f"输出文件: {output_path}")

def main():
    parser = argparse.ArgumentParser(description='AUTOSAR NvM 配置同步工具 (最终版)')
    
    parser.add_argument('-i', '--input', required=True, help='输入的 ARXML 文件路径')
    parser.add_argument('-e', '--excel', required=True, help='输入的 Excel 文件路径')
    parser.add_argument('-o', '--output', help='输出的 ARXML 文件路径 (可选)')
    
    args = parser.parse_args()
    
    if args.output:
        output_file = args.output
    else:
        base_name = os.path.splitext(args.input)[0]
        ext = os.path.splitext(args.input)[1]
        output_file = f"{base_name}_Updated{ext}"
    
    if not os.path.exists(args.input):
        print(f"错误: 找不到输入文件 {args.input}")
        sys.exit(1)
        
    if not os.path.exists(args.excel):
        print(f"错误: 找不到 Excel 文件 {args.excel}")
        sys.exit(1)

    try:
        print(f"正在读取 Excel: {args.excel} ...")
        excel_map = load_excel_data(args.excel)
        print(f"从 Excel 读取到 {len(excel_map)} 条记录。")
        
        print(f"正在处理 ARXML: {args.input} ...")
        update_xml(args.input, excel_map, output_file)
        
    except Exception as e:
        import traceback
        print(f"发生严重错误: {e}")
        traceback.print_exc()
        sys.exit(1)

if __name__ == '__main__':
    main()